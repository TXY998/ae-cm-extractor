import os
import streamlit as st
import pandas as pd
import json
import openai
from io import BytesIO
from datetime import datetime
from docx import Document

# ==================== 页面配置 ====================
st.set_page_config(
    page_title="AE/CM 智能提取系统",
    page_icon="🏥",
    layout="wide",
    initial_sidebar_state="expanded"
)

# ==================== 自定义CSS ====================
st.markdown("""
<style>
    .main-header {
        font-size: 2rem;
        font-weight: 700;
        color: #1f3a5f;
        border-bottom: 3px solid #1f3a5f;
        padding-bottom: 10px;
        margin-bottom: 20px;
    }
    .sub-header {
        font-size: 1.3rem;
        font-weight: 600;
        color: #2d5a87;
        margin-top: 20px;
    }
    .design-note {
        background-color: #e8f4f8;
        border-left: 4px solid #2d5a87;
        padding: 12px 16px;
        border-radius: 4px;
        margin: 16px 0;
    }
</style>
""", unsafe_allow_html=True)

# ==================== 初始化session_state ====================
if "ae_data" not in st.session_state:
    st.session_state["ae_data"] = []
if "cm_data" not in st.session_state:
    st.session_state["cm_data"] = []
if "original_text" not in st.session_state:
    st.session_state["original_text"] = ""
if "extraction_done" not in st.session_state:
    st.session_state["extraction_done"] = False
if "generated_text" not in st.session_state:
    st.session_state["generated_text"] = ""
if "gen_ae" not in st.session_state:
    st.session_state["gen_ae"] = []
if "gen_cm" not in st.session_state:
    st.session_state["gen_cm"] = []

# ==================== 侧边栏 ====================
with st.sidebar:
    st.markdown("## ⚙️ 系统配置")

    # 从环境变量读取预置的API Key（如果存在）
    default_api_key = os.environ.get("ZHIPU_API_KEY", "")

    api_key = st.text_input(
        "🔑 API Key",
        type="password",
        value=default_api_key,
        placeholder="已预填演示Key，可直接使用" if default_api_key else "请输入智谱AI API Key",
        help="已自动填入演示Key，直接点击提取即可。如无效，请前往 open.bigmodel.cn 免费获取（1分钟注册）。"
    )

    st.markdown("[📌 获取API Key（新用户免费）](https://open.bigmodel.cn)")

    st.divider()

    st.markdown("### 📖 CTCAE v5.0 快速参考")
    with st.expander("点击展开分级标准"):
        ctcae_ref = pd.DataFrame({
            "等级": ["Grade 1", "Grade 2", "Grade 3", "Grade 4", "Grade 5"],
            "严重程度": ["轻度", "中度", "重度", "危及生命", "死亡"],
            "描述": [
                "无症状或轻微；仅临床或诊断观察到",
                "需要局部或非侵入性干预",
                "严重但不立即危及生命；需住院或延长住院",
                "危及生命；需要紧急干预",
                "与AE相关的死亡"
            ]
        })
        st.dataframe(ctcae_ref, hide_index=True, use_container_width=True)

    st.divider()

    st.markdown("### 💡 设计理念")
    st.markdown("""
    <div class="design-note">
    <strong>AI 初筛 + 原文证据回溯 + 人工确认</strong><br><br>
    <strong>系统定位：辅助工具，非替代医学判断</strong><br>
    • 大模型负责初筛，降低人工成本<br>
    • 每条记录保留原文证据，可溯源<br>
    • 双向转换：病历↔AE/CM表格<br>
    • 所有结果需人工确认后使用
    </div>
    """, unsafe_allow_html=True)

    st.caption(f"© 2026 AE/CM Extractor v2.0 | {datetime.now().strftime('%Y-%m-%d')}")

# ==================== 标题区 ====================
st.markdown('<p class="main-header">🏥 AE/CM 智能提取与核验系统</p>', unsafe_allow_html=True)
st.caption("临床试验不良事件与合并用药的 AI 辅助提取工具 | 支持病历↔表格双向转换")

# ==================== 主界面：5个标签页 ====================
tab_extract, tab_ae, tab_cm, tab_generate, tab_export = st.tabs([
    "📝 1. 病历→AE/CM",
    "🔴 2. AE 不良事件表",
    "💊 3. CM 合并用药表",
    "🔄 4. AE/CM→病历",
    "📥 5. 导出与复核"
])

# ============================================================
# 标签页 1：病历 → AE/CM
# ============================================================
with tab_extract:
    st.markdown('<p class="sub-header">📄 输入病历/随访文本，自动提取AE和CM</p>', unsafe_allow_html=True)

    col1, col2 = st.columns([3, 1])
    with col1:
        st.caption("粘贴出院小结、随访记录，AI自动识别并提取不良事件和合并用药")
    with col2:
        show_example = st.checkbox("📋 使用示例文本", value=True)

    example_text = """患者，男，65岁。2026.03.01开始出现恶心，03.03呕吐2次，予甲氧氯普胺10mg tid口服治疗后好转；
03.05复查血常规：白细胞2.1×10⁹/L，考虑白细胞减少，予升白针（重组人粒细胞集落刺激因子）150μg皮下注射治疗；
03.08恶心呕吐症状完全消失，停用甲氧氯普胺；
03.10复查白细胞升至4.5×10⁹/L，继续观察。"""

    input_text = st.text_area(
        "病历文本",
        value=example_text if show_example else "",
        height=220,
        placeholder="请粘贴病历/随访文本...",
        key="extract_input"
    )

    col_btn1, col_btn2 = st.columns([1, 4])
    with col_btn1:
        extract_btn = st.button("🚀 提取 AE/CM", type="primary", use_container_width=True)
    with col_btn2:
        if st.button("🔄 清空所有数据", use_container_width=True):
            st.session_state["ae_data"] = []
            st.session_state["cm_data"] = []
            st.session_state["original_text"] = ""
            st.session_state["extraction_done"] = False
            st.session_state["generated_text"] = ""
            st.session_state["gen_ae"] = []
            st.session_state["gen_cm"] = []
            st.rerun()

    if extract_btn:
        if not api_key:
            st.error("❌ 请先在左侧填写 API Key！")
        elif not input_text.strip():
            st.warning("⚠️ 请输入病历文本后再提取。")
        else:
            with st.spinner("🤖 AI 正在分析文本，提取AE/CM信息..."):
                try:
                    client = openai.OpenAI(
                        api_key=api_key,
                        base_url="https://open.bigmodel.cn/api/paas/v4/"
                    )

                    extraction_prompt = f"""
你是一名有10年经验的临床试验数据管理专家。
请从以下病历文本中提取**所有**不良事件(AE)和合并用药(CM)信息，**不能遗漏任何一条**。

=== 提取原则 ===
1. 只提取文本中明确提及的信息，绝不编造。
2. 每条记录必须包含"原文证据"——即文本中支持该提取的原句或短语。
3. AE包括但不限于：临床症状(恶心、呕吐、疼痛等)、体征、实验室检查异常(白细胞减少、血小板降低、肝酶升高等)。
4. CTCAE术语建议使用标准英文术语，等级建议填入数字1-5；若无法确定填"待确认"。
5. "与试验药物相关性"根据原文推断，可选值：肯定有关、可能有关、可能无关、无关、待确认。若文本未提及或无法推断，填"待确认"。
6. 若字段文本未提及，填"未提及"。

=== 特别注意 ===
- 文中明确写了"恶心"、"呕吐"，必须分别提取为独立的AE。
- 文中明确写了"白细胞2.1×10⁹/L，考虑白细胞减少"，必须提取为白细胞减少。
- 每种药物/治疗措施都必须提取到CM中。

=== 输出格式 ===
严格返回以下JSON，不要包含其他任何内容：
{{
  "ae": [
    {{
      "id": 1,
      "ae_name": "恶心",
      "evidence": "03.01开始出现恶心",
      "start_date": "2026.03.01",
      "end_date": "2026.03.08",
      "ctcae_term": "Nausea",
      "grade": "待确认",
      "relationship": "待确认",
      "status": "待人工确认"
    }},
    {{
      "id": 2,
      "ae_name": "呕吐",
      "evidence": "03.03呕吐2次",
      "start_date": "2026.03.03",
      "end_date": "2026.03.08",
      "ctcae_term": "Vomiting",
      "grade": "待确认",
      "relationship": "待确认",
      "status": "待人工确认"
    }},
    {{
      "id": 3,
      "ae_name": "白细胞减少",
      "evidence": "白细胞2.1×10⁹/L，考虑白细胞减少",
      "start_date": "2026.03.05",
      "end_date": "未提及",
      "ctcae_term": "White blood cell decreased",
      "grade": "待确认",
      "relationship": "待确认",
      "status": "待人工确认"
    }}
  ],
  "cm": [
    {{
      "id": 1,
      "drug_name": "甲氧氯普胺",
      "related_ae": "恶心/呕吐",
      "start_date": "2026.03.03",
      "end_date": "2026.03.08",
      "dose": "10mg",
      "frequency": "tid",
      "route": "口服",
      "evidence": "予甲氧氯普胺10mg tid口服治疗",
      "status": "待人工确认"
    }},
    {{
      "id": 2,
      "drug_name": "重组人粒细胞集落刺激因子（升白针）",
      "related_ae": "白细胞减少",
      "start_date": "2026.03.05",
      "end_date": "未提及",
      "dose": "150μg",
      "frequency": "未提及",
      "route": "皮下注射",
      "evidence": "予升白针150μg皮下注射治疗",
      "status": "待人工确认"
    }}
  ]
}}

文本：
{input_text}
"""

                    response = client.chat.completions.create(
                        model="glm-4-flash",
                        messages=[{"role": "user", "content": extraction_prompt}],
                        temperature=0.1
                    )

                    result_text = response.choices[0].message.content.strip()

                    if "```json" in result_text:
                        result_text = result_text.split("```json")[1]
                    if "```" in result_text:
                        result_text = result_text.split("```")[0]
                    result_text = result_text.strip()

                    data = json.loads(result_text)

                    st.session_state["ae_data"] = data.get("ae", [])
                    st.session_state["cm_data"] = data.get("cm", [])
                    st.session_state["original_text"] = input_text
                    st.session_state["extraction_done"] = True

                    ae_count = len(st.session_state["ae_data"])
                    cm_count = len(st.session_state["cm_data"])

                    st.success(f"✅ 提取完成！共识别 {ae_count} 条不良事件，{cm_count} 条合并用药。")
                    st.info("💡 请切换到「AE不良事件表」和「CM合并用药表」标签页查看详情并人工确认。")

                except json.JSONDecodeError as e:
                    st.error(f"❌ AI返回格式解析失败，请重试。\n错误：{str(e)}")
                except Exception as e:
                    st.error(f"❌ 提取失败：{str(e)}")

# ============================================================
# 标签页 2：AE 表
# ============================================================
with tab_ae:
    st.markdown('<p class="sub-header">🔴 不良事件（AE）列表</p>', unsafe_allow_html=True)

    if not st.session_state["extraction_done"]:
        st.info("📭 请先在「病历→AE/CM」标签页中提取数据。")
    else:
        st.caption("✏️ 双击单元格修改 | 等级填数字1-5 | 状态可批量填充")

        ae_df = pd.DataFrame(st.session_state["ae_data"])
        columns_order = ["id", "ae_name", "evidence", "start_date", "end_date", "ctcae_term", "grade", "relationship", "status"]
        for col in columns_order:
            if col not in ae_df.columns:
                ae_df[col] = ""
        ae_df = ae_df[columns_order]

        if len(ae_df) == 0:
            st.info("📭 未提取到AE数据。可在下方手动添加。")
            ae_df = pd.DataFrame(columns=columns_order)

        st.markdown("""
        <div class="design-note">
        🔍 <strong>原文证据回溯：</strong>每条记录的"原文证据"列显示AI提取依据的原始文本，可直接溯源验证。
        等级请填写数字 <strong>1、2、3、4、5</strong>，不确定则填 <strong>"待确认"</strong>。
        </div>
        """, unsafe_allow_html=True)

        edited_ae = st.data_editor(
            ae_df,
            num_rows="dynamic",
            use_container_width=True,
            hide_index=True,
            key="ae_editor",
            column_config={
                "id": st.column_config.TextColumn("序号", width="small", disabled=True),
                "ae_name": st.column_config.TextColumn("AE名称", width="medium"),
                "evidence": st.column_config.TextColumn("📋 原文证据", width="large"),
                "start_date": st.column_config.TextColumn("开始日期", width="medium"),
                "end_date": st.column_config.TextColumn("结束日期", width="medium"),
                "ctcae_term": st.column_config.TextColumn("CTCAE术语建议", width="medium"),
                "grade": st.column_config.TextColumn(
                    "等级(1-5)",
                    width="small",
                    help="CTCAE等级：1=轻度 2=中度 3=重度 4=危及生命 5=死亡"
                ),
                "relationship": st.column_config.TextColumn(
                    "与试验药物相关性",
                    width="medium",
                    help="可选：肯定有关、可能有关、可能无关、无关、待确认"
                ),
                "status": st.column_config.TextColumn(
                    "状态",
                    width="small",
                    help="待人工确认 / ✅已确认 / ✏️已修改 / ❌已删除"
                )
            }
        )

        st.caption("🎯 等级快捷填充：")
        cols_grade = st.columns(6)
        grade_map = [
            ("1级", "1"),
            ("2级", "2"),
            ("3级", "3"),
            ("4级", "4"),
            ("5级", "5"),
            ("待确认", "待确认")
        ]
        for i, (label, val) in enumerate(grade_map):
            with cols_grade[i]:
                if st.button(label, use_container_width=True, key=f"ae_grade_{val}"):
                    for idx in range(len(edited_ae)):
                        edited_ae.at[idx, "grade"] = val
                    st.session_state["ae_data"] = edited_ae.to_dict("records")
                    st.rerun()

        st.caption("🔗 相关性快捷填充：")
        cols_rel = st.columns(5)
        rel_options = ["肯定有关", "可能有关", "可能无关", "无关", "待确认"]
        for i, rel_val in enumerate(rel_options):
            with cols_rel[i]:
                if st.button(rel_val, use_container_width=True, key=f"ae_rel_{rel_val}"):
                    for idx in range(len(edited_ae)):
                        edited_ae.at[idx, "relationship"] = rel_val
                    st.session_state["ae_data"] = edited_ae.to_dict("records")
                    st.rerun()

        st.caption("💡 状态快捷填充：")
        cols_status = st.columns(5)
        status_options = ["待人工确认", "✅已确认", "✏️已修改", "❌已删除"]
        for i, status_val in enumerate(status_options):
            with cols_status[i]:
                if st.button(status_val, use_container_width=True, key=f"ae_status_{status_val}"):
                    for idx in range(len(edited_ae)):
                        edited_ae.at[idx, "status"] = status_val
                    st.session_state["ae_data"] = edited_ae.to_dict("records")
                    st.rerun()

        st.divider()

        ae_total = len(edited_ae)
        ae_confirmed = sum(1 for _, row in edited_ae.iterrows() if "已确认" in str(row.get("status", "")))
        ae_graded = sum(1 for _, row in edited_ae.iterrows() if str(row.get("grade", "")).isdigit())

        col_m1, col_m2, col_m3 = st.columns(3)
        with col_m1:
            st.metric("AE总数", ae_total)
        with col_m2:
            st.metric("已确认", ae_confirmed)
        with col_m3:
            st.metric("已分级", ae_graded)

        if st.button("💾 保存AE修改", use_container_width=True, type="primary", key="ae_save"):
            st.session_state["ae_data"] = edited_ae.to_dict("records")
            st.success("✅ AE表修改已保存")

        st.session_state["ae_data"] = edited_ae.to_dict("records")

# ============================================================
# 标签页 3：CM 表
# ============================================================
with tab_cm:
    st.markdown('<p class="sub-header">💊 合并用药（CM）列表</p>', unsafe_allow_html=True)

    if not st.session_state["extraction_done"]:
        st.info("📭 请先在「病历→AE/CM」标签页中提取数据。")
    else:
        st.caption("✏️ 双击单元格修改 | 剂量/频次/途径各自独立")

        cm_df = pd.DataFrame(st.session_state["cm_data"])
        columns_order_cm = ["id", "drug_name", "related_ae", "start_date", "end_date", "dose", "frequency", "route", "evidence", "status"]
        for col in columns_order_cm:
            if col not in cm_df.columns:
                cm_df[col] = ""
        cm_df = cm_df[columns_order_cm]

        if len(cm_df) == 0:
            st.info("📭 未提取到CM数据。可在下方手动添加。")
            cm_df = pd.DataFrame(columns=columns_order_cm)

        st.markdown("""
        <div class="design-note">
        🔗 <strong>AE-CM关联：</strong>每条CM记录关联对应的AE事件，便于追溯用药原因。
        剂量、频次、途径已拆分为独立列。
        </div>
        """, unsafe_allow_html=True)

        edited_cm = st.data_editor(
            cm_df,
            num_rows="dynamic",
            use_container_width=True,
            hide_index=True,
            key="cm_editor",
            column_config={
                "id": st.column_config.TextColumn("序号", width="small", disabled=True),
                "drug_name": st.column_config.TextColumn("药品名称", width="medium"),
                "related_ae": st.column_config.TextColumn("🔗 对应事件", width="medium"),
                "start_date": st.column_config.TextColumn("开始日期", width="medium"),
                "end_date": st.column_config.TextColumn("结束日期", width="medium"),
                "dose": st.column_config.TextColumn("剂量", width="small"),
                "frequency": st.column_config.TextColumn("频次", width="small"),
                "route": st.column_config.TextColumn("途径", width="small"),
                "evidence": st.column_config.TextColumn("📋 原文证据", width="large"),
                "status": st.column_config.TextColumn(
                    "状态",
                    width="small",
                    help="待人工确认 / ✅已确认 / ✏️已修改 / ❌已删除"
                )
            }
        )

        st.caption("💊 途径快捷填充：")
        cols_route = st.columns(6)
        route_options = ["口服", "静脉注射", "肌注", "皮下注射", "外用", "未提及"]
        for i, route_val in enumerate(route_options):
            with cols_route[i]:
                if st.button(route_val, use_container_width=True, key=f"cm_route_{route_val}"):
                    for idx in range(len(edited_cm)):
                        edited_cm.at[idx, "route"] = route_val
                    st.session_state["cm_data"] = edited_cm.to_dict("records")
                    st.rerun()

        st.caption("💡 状态快捷填充：")
        cols_status_cm = st.columns(5)
        for i, status_val in enumerate(status_options):
            with cols_status_cm[i]:
                if st.button(status_val, use_container_width=True, key=f"cm_status_{status_val}"):
                    for idx in range(len(edited_cm)):
                        edited_cm.at[idx, "status"] = status_val
                    st.session_state["cm_data"] = edited_cm.to_dict("records")
                    st.rerun()

        st.divider()

        cm_total = len(edited_cm)
        cm_confirmed = sum(1 for _, row in edited_cm.iterrows() if "已确认" in str(row.get("status", "")))
        cm_with_route = sum(1 for _, row in edited_cm.iterrows() if str(row.get("route", "")) not in ["", "未提及"])

        col_c1, col_c2, col_c3 = st.columns(3)
        with col_c1:
            st.metric("CM总数", cm_total)
        with col_c2:
            st.metric("已确认", cm_confirmed)
        with col_c3:
            st.metric("已标注途径", cm_with_route)

        if st.button("💾 保存CM修改", use_container_width=True, type="primary", key="cm_save"):
            st.session_state["cm_data"] = edited_cm.to_dict("records")
            st.success("✅ CM表修改已保存")

        st.session_state["cm_data"] = edited_cm.to_dict("records")

# ============================================================
# 标签页 4：AE/CM → 病历（反向生成）
# ============================================================
with tab_generate:
    st.markdown('<p class="sub-header">🔄 从AE/CM表格反向生成病历文本</p>', unsafe_allow_html=True)
    st.caption("在下表中录入AE和CM数据，AI将自动生成结构化的病历/随访记录文本")

    if st.session_state["extraction_done"] and (st.session_state["ae_data"] or st.session_state["cm_data"]):
        with st.expander("📋 从已提取的数据导入（可选）", expanded=False):
            if st.button("📥 一键导入已提取的AE/CM数据", use_container_width=True):
                st.session_state["gen_ae"] = st.session_state["ae_data"]
                st.session_state["gen_cm"] = st.session_state["cm_data"]
                st.success("✅ 已导入！")
                st.rerun()

    st.markdown("#### 🔴 不良事件（AE）录入")
    gen_ae_df = pd.DataFrame(st.session_state["gen_ae"]) if st.session_state["gen_ae"] else pd.DataFrame(
        columns=["ae_name", "grade", "start_date", "end_date", "outcome", "relationship"]
    )
    for col in ["ae_name", "grade", "start_date", "end_date", "outcome", "relationship"]:
        if col not in gen_ae_df.columns:
            gen_ae_df[col] = ""

    edited_gen_ae = st.data_editor(
        gen_ae_df,
        num_rows="dynamic",
        use_container_width=True,
        hide_index=True,
        key="gen_ae_editor",
        column_config={
            "ae_name": st.column_config.TextColumn("AE名称*", width="medium"),
            "grade": st.column_config.TextColumn("等级(1-5)", width="small"),
            "start_date": st.column_config.TextColumn("开始日期", width="medium"),
            "end_date": st.column_config.TextColumn("结束日期", width="medium"),
            "outcome": st.column_config.TextColumn("结局", width="medium"),
            "relationship": st.column_config.TextColumn("与药物关系", width="medium")
        }
    )

    st.markdown("#### 💊 合并用药（CM）录入")
    gen_cm_df = pd.DataFrame(st.session_state["gen_cm"]) if st.session_state["gen_cm"] else pd.DataFrame(
        columns=["drug_name", "dose", "frequency", "route", "start_date", "end_date", "related_ae", "indication"]
    )
    for col in ["drug_name", "dose", "frequency", "route", "start_date", "end_date", "related_ae", "indication"]:
        if col not in gen_cm_df.columns:
            gen_cm_df[col] = ""

    edited_gen_cm = st.data_editor(
        gen_cm_df,
        num_rows="dynamic",
        use_container_width=True,
        hide_index=True,
        key="gen_cm_editor",
        column_config={
            "drug_name": st.column_config.TextColumn("药品名称*", width="medium"),
            "dose": st.column_config.TextColumn("剂量", width="small"),
            "frequency": st.column_config.TextColumn("频次", width="small"),
            "route": st.column_config.TextColumn("途径", width="small"),
            "start_date": st.column_config.TextColumn("开始日期", width="medium"),
            "end_date": st.column_config.TextColumn("结束日期", width="medium"),
            "related_ae": st.column_config.TextColumn("对应AE", width="medium"),
            "indication": st.column_config.TextColumn("用药原因", width="medium")
        }
    )

    with st.expander("📋 患者基本信息（可选）"):
        col1, col2, col3 = st.columns(3)
        with col1:
            patient_info = st.text_input("姓名/编号", value="患者", key="gen_patient")
            patient_gender = st.selectbox("性别", ["男", "女", "未指定"], key="gen_gender")
        with col2:
            patient_age = st.text_input("年龄", value="", key="gen_age")
            patient_diagnosis = st.text_input("诊断", value="", key="gen_diagnosis")
        with col3:
            trial_drug = st.text_input("试验药物", value="", key="gen_drug")
            visit_date = st.text_input("访视日期", value="", key="gen_date")

    st.divider()
    gen_btn = st.button("🔄 反向生成病历文本", type="primary", use_container_width=True, key="generate_btn")

    if gen_btn:
        if not api_key:
            st.error("❌ 请先在左侧填写 API Key！")
        else:
            has_ae = any(pd.notna(row.get("ae_name")) and str(row["ae_name"]).strip() for _, row in edited_gen_ae.iterrows())
            has_cm = any(pd.notna(row.get("drug_name")) and str(row["drug_name"]).strip() for _, row in edited_gen_cm.iterrows())

            if not has_ae and not has_cm:
                st.warning("⚠️ 请至少录入一条AE或CM数据。")
            else:
                with st.spinner("🤖 AI 正在根据AE/CM数据生成病历文本..."):
                    try:
                        client = openai.OpenAI(
                            api_key=api_key,
                            base_url="https://open.bigmodel.cn/api/paas/v4/"
                        )

                        ae_list = []
                        for _, row in edited_gen_ae.iterrows():
                            if pd.notna(row.get("ae_name")) and str(row["ae_name"]).strip():
                                ae_list.append({
                                    "名称": str(row["ae_name"]),
                                    "等级": str(row.get("grade", "未提及")),
                                    "开始日期": str(row.get("start_date", "未提及")),
                                    "结束日期": str(row.get("end_date", "未提及")),
                                    "结局": str(row.get("outcome", "未提及")),
                                    "与药物关系": str(row.get("relationship", "未提及"))
                                })

                        cm_list = []
                        for _, row in edited_gen_cm.iterrows():
                            if pd.notna(row.get("drug_name")) and str(row["drug_name"]).strip():
                                cm_list.append({
                                    "药品名称": str(row["drug_name"]),
                                    "剂量": str(row.get("dose", "未提及")),
                                    "频次": str(row.get("frequency", "未提及")),
                                    "途径": str(row.get("route", "未提及")),
                                    "开始日期": str(row.get("start_date", "未提及")),
                                    "结束日期": str(row.get("end_date", "未提及")),
                                    "对应AE": str(row.get("related_ae", "未提及")),
                                    "用药原因": str(row.get("indication", "未提及"))
                                })

                        ae_json = json.dumps(ae_list, ensure_ascii=False, indent=2)
                        cm_json = json.dumps(cm_list, ensure_ascii=False, indent=2)

                        generate_prompt = f"""
你是一名经验丰富的临床研究医生。请根据以下不良事件（AE）和合并用药（CM）信息，撰写一段结构化的病历/随访记录文本。

患者信息：
- 姓名/编号：{patient_info}
- 性别：{patient_gender}
- 年龄：{patient_age if patient_age else "未提供"}
- 诊断：{patient_diagnosis if patient_diagnosis else "未提供"}
- 试验药物：{trial_drug if trial_drug else "未提供"}
- 访视日期：{visit_date if visit_date else "未提供"}

不良事件（AE）：
{ae_json if ae_list else "（无不良事件报告）"}

合并用药（CM）：
{cm_json if cm_list else "（无合并用药）"}

要求：
1. 用自然流畅的临床叙述语言，按时间顺序组织内容
2. 包含患者基本信息、用药情况、不良事件描述
3. 语言专业但不过度生硬，像一份真实的病历记录
4. 直接输出最终文本，不要加任何引导语或解释
5. 如果剂量/频次/途径信息完整，请按标准格式书写
"""

                        response = client.chat.completions.create(
                            model="glm-4-flash",
                            messages=[{"role": "user", "content": generate_prompt}],
                            temperature=0.7
                        )

                        st.session_state["generated_text"] = response.choices[0].message.content.strip()
                        st.session_state["gen_ae"] = edited_gen_ae.to_dict("records")
                        st.session_state["gen_cm"] = edited_gen_cm.to_dict("records")

                        st.success("✅ 病历文本生成完成！")

                    except Exception as e:
                        st.error(f"❌ 生成失败：{str(e)}")

    if st.session_state["generated_text"]:
        st.divider()
        st.markdown("#### 📝 生成的病历文本")
        st.caption("✏️ 可编辑、可复制、可导出")

        edited_generated = st.text_area(
            "病历文本",
            value=st.session_state["generated_text"],
            height=300,
            key="generated_text_display"
        )

        col_down1, col_down2 = st.columns(2)
        with col_down1:
            st.download_button(
                label="📥 导出为 TXT",
                data=edited_generated.encode("utf-8"),
                file_name=f"生成的病历_{datetime.now().strftime('%Y%m%d_%H%M')}.txt",
                mime="text/plain"
            )
        with col_down2:
            output_word = BytesIO()
            doc = Document()
            doc.add_heading("病历/随访记录", level=1)
            for line in edited_generated.split("\n"):
                if line.strip():
                    doc.add_paragraph(line)
            doc.save(output_word)
            output_word.seek(0)
            st.download_button(
                label="📥 导出为 Word",
                data=output_word,
                file_name=f"生成的病历_{datetime.now().strftime('%Y%m%d_%H%M')}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )

# ============================================================
# 标签页 5：导出与复核
# ============================================================
with tab_export:
    st.markdown('<p class="sub-header">📥 导出与复核报告</p>', unsafe_allow_html=True)

    if not st.session_state["extraction_done"]:
        st.info("📭 请先在「病历→AE/CM」标签页中提取数据。")
    else:
        ae_export = pd.DataFrame(st.session_state["ae_data"]) if st.session_state["ae_data"] else pd.DataFrame()
        cm_export = pd.DataFrame(st.session_state["cm_data"]) if st.session_state["cm_data"] else pd.DataFrame()

        ae_total = len(ae_export)
        ae_confirmed = 0
        if not ae_export.empty and "status" in ae_export.columns:
            ae_confirmed = sum(1 for _, row in ae_export.iterrows() if "已确认" in str(row.get("status", "")))

        cm_total = len(cm_export)
        cm_confirmed = 0
        if not cm_export.empty and "status" in cm_export.columns:
            cm_confirmed = sum(1 for _, row in cm_export.iterrows() if "已确认" in str(row.get("status", "")))

        st.markdown("### 📊 复核统计")
        col_stat1, col_stat2, col_stat3, col_stat4 = st.columns(4)
        with col_stat1:
            st.metric("AE总数", ae_total)
        with col_stat2:
            st.metric("AE已确认", ae_confirmed)
        with col_stat3:
            st.metric("CM总数", cm_total)
        with col_stat4:
            st.metric("CM已确认", cm_confirmed)

        st.divider()
        st.markdown("### 📦 导出文件")

        col_exp1, col_exp2, col_exp3 = st.columns(3)

        with col_exp1:
            st.markdown("#### 🔴 AE 表")
            output_ae = BytesIO()
            with pd.ExcelWriter(output_ae, engine="openpyxl") as writer:
                ae_export.to_excel(writer, sheet_name="不良事件AE", index=False)
            st.download_button(
                label="⬇️ 下载 AE表.xlsx",
                data=output_ae.getvalue(),
                file_name=f"AE_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

        with col_exp2:
            st.markdown("#### 💊 CM 表")
            output_cm = BytesIO()
            with pd.ExcelWriter(output_cm, engine="openpyxl") as writer:
                cm_export.to_excel(writer, sheet_name="合并用药CM", index=False)
            st.download_button(
                label="⬇️ 下载 CM表.xlsx",
                data=output_cm.getvalue(),
                file_name=f"CM_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

        with col_exp3:
            st.markdown("#### 📋 核对报告")
            output_report = BytesIO()
            with pd.ExcelWriter(output_report, engine="openpyxl") as writer:
                summary_data = {
                    "项目": ["提取时间", "AE数量", "AE已确认", "CM数量", "CM已确认"],
                    "内容": [
                        datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
                        str(ae_total), str(ae_confirmed),
                        str(cm_total), str(cm_confirmed)
                    ]
                }
                pd.DataFrame(summary_data).to_excel(writer, sheet_name="核对摘要", index=False)
                if not ae_export.empty:
                    ae_export.to_excel(writer, sheet_name="AE详情", index=False)
                if not cm_export.empty:
                    cm_export.to_excel(writer, sheet_name="CM详情", index=False)
                pd.DataFrame({"原始病历文本": [st.session_state["original_text"]]}).to_excel(
                    writer, sheet_name="原始文本", index=False
                )
            st.download_button(
                label="⬇️ 下载 核对报告.xlsx",
                data=output_report.getvalue(),
                file_name=f"核对报告_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

        st.divider()
        st.markdown("### 👁️ 数据预览")
        tab_p1, tab_p2 = st.tabs(["AE 预览", "CM 预览"])
        with tab_p1:
            if not ae_export.empty:
                st.dataframe(ae_export, use_container_width=True, hide_index=True)
            else:
                st.info("无AE数据")
        with tab_p2:
            if not cm_export.empty:
                st.dataframe(cm_export, use_container_width=True, hide_index=True)
            else:
                st.info("无CM数据")

# ==================== 页脚 ====================
st.divider()
st.caption("⚠️ 免责声明：本系统为AI辅助工具，所有输出结果需经专业医学人员审核确认后方可使用。")
st.caption(f"🏥 AE/CM 智能提取与核验系统 v2.0 | Powered by 智谱GLM-4 | {datetime.now().year}")
